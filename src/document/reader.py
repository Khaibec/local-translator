"""Document reader abstraction and implementations."""

from abc import ABC, abstractmethod
from pathlib import Path
from typing import Type
from src.utils.text_utils import normalize_line_endings


class BaseDocumentReader(ABC):
    """Abstract base class for document readers."""

    @abstractmethod
    def read(self, file_path: Path) -> str:
        """Read and return document content as a string."""
        pass


class TxtReader(BaseDocumentReader):
    """Reader for plain text files with encoding detection for Japanese text."""

    def read(self, file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"Source file not found: {file_path}")

        # Try UTF-8-sig first (handles BOM cleanly), then standard UTF-8, then Japanese encodings
        encodings = ["utf-8-sig", "utf-8", "cp932", "shift_jis", "euc-jp"]
        content = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if content is None:
            raise UnicodeDecodeError(
                "utf-8", b"", 0, 0,
                f"Failed to decode '{file_path}' using supported encodings ({', '.join(encodings)})."
            )

        return normalize_line_endings(content)


class DocxReader(BaseDocumentReader):
    """Reader for Word (.docx) documents."""

    def read(self, file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"Source file not found: {file_path}")

        try:
            import docx
        except ImportError:
            raise ImportError("python-docx is required to read DOCX files. Install it with: pip install python-docx")

        doc = docx.Document(file_path)
        paragraphs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        # Also read table cells if any
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))

        return normalize_line_endings("\n\n".join(paragraphs))


class PdfReader(BaseDocumentReader):
    """Reader for text-based PDF documents."""

    def read(self, file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"Source file not found: {file_path}")

        try:
            import pypdf
        except ImportError:
            raise ImportError("pypdf is required to read PDF files. Install it with: pip install pypdf")

        pages_text = []
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted and extracted.strip():
                    pages_text.append(extracted.strip())

        return normalize_line_endings("\n\n".join(pages_text))


def get_reader(file_path: Path) -> BaseDocumentReader:
    """Factory to retrieve appropriate document reader based on file suffix."""
    suffix = file_path.suffix.lower()
    if suffix in [".txt", ".text", ".md"]:
        return TxtReader()
    elif suffix == ".docx":
        return DocxReader()
    elif suffix == ".pdf":
        return PdfReader()
    else:
        raise ValueError(f"Unsupported file format: '{suffix}'. Supported formats: .txt, .md, .docx, .pdf")
